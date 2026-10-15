"""Unit and Integration Tests for Category-wise Attendance Storage & Document Downloads.
Tests:
1. Category-wise attendance filtering (Staff, Employees, Contract Employees).
2. Master Excel workbook generation with Day 1 to 31, Total Days Come, and Contract OT sheet.
3. Individual Employee Attendance Document named {emp_code}_attendance_{month}.xlsx:
   - Day 1 to 31 rows
   - Total Days Come (e.g. 30 days)
   - Contract OT breakdown (10 hrs present + 2 hrs OT)
4. HR and Employee download endpoints (/hr/reports/employee/<eid>/export, /history/export).
"""
import io
import os
import shutil
import tempfile
import unittest
from datetime import date
import openpyxl

from app import create_app
from app.db import get_db, init_db
from app.services import d2s, today
from app.exports import build_employee_attendance_doc, build_workbook


class CategoryAndDocsTestCase(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmpdir = tempfile.mkdtemp()
        cls.db_path = os.path.join(cls.tmpdir, "test_cat_docs.db")
        cls.selfie_dir = os.path.join(cls.tmpdir, "selfies")
        os.makedirs(cls.selfie_dir, exist_ok=True)

        os.environ["DATABASE"] = cls.db_path
        os.environ["SELFIE_DIR"] = cls.selfie_dir
        os.environ["APP_ENV"] = "testing"
        os.environ["DEMO_DATA"] = "0"

        cls.app = create_app()
        cls.app.config["TESTING"] = True
        cls.app.config["DATABASE"] = cls.db_path
        cls.app.config["SELFIE_DIR"] = cls.selfie_dir
        cls.app.config["DEMO_DATA"] = False

        init_db(cls.app)
        cls.client = cls.app.test_client()

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmpdir, ignore_errors=True)

    def setUp(self):
        with self.app.app_context():
            db = get_db()
            # Ensure categories exist
            contract_cat = db.execute("SELECT id FROM employee_categories WHERE code='CONTRACT'").fetchone()[0]
            staff_cat = db.execute("SELECT id FROM employee_categories WHERE code='STAFF'").fetchone()[0]
            s1 = db.execute("SELECT id FROM shifts ORDER BY id LIMIT 1").fetchone()[0]

            # Insert a contract employee CON001
            existing = db.execute("SELECT id FROM employees WHERE emp_code='CON001'").fetchone()
            if not existing:
                db.execute(
                    "INSERT INTO employees (emp_code, name, email, department, designation, role, shift_id, office_id, category_id, password_hash, must_change_pw, joined_on) "
                    "VALUES ('CON001', 'Contract Worker Ram', 'ram@sbg.com', 'Assembly', 'Technician', 'employee', ?, 1, ?, 'hash', 0, '2026-01-01')",
                    (s1, contract_cat)
                )
                db.commit()

            con_emp = db.execute("SELECT * FROM employees WHERE emp_code='CON001'").fetchone()

            # Seed 30 days of attendance for CON001 in October 2026: 10 hrs regular + 2 hrs OT = 12 hrs
            for day in range(1, 31):
                d_str = f"2026-10-{day:02d}"
                db.execute(
                    "INSERT OR REPLACE INTO attendance (emp_id, date, shift_id, check_in, check_out, work_hours, ot_hours, status, late) "
                    "VALUES (?,?,?,?,?,?,?,?,0)",
                    (con_emp["id"], d_str, s1, f"{d_str} 08:00:00", f"{d_str} 20:00:00", 12.0, 2.0, "OT")
                )
            # Day 31: Absent
            db.execute(
                "INSERT OR REPLACE INTO attendance (emp_id, date, shift_id, check_in, check_out, work_hours, ot_hours, status, late) "
                "VALUES (?,?,?,NULL,NULL,0.0,0.0,'A',0)",
                (con_emp["id"], "2026-10-31", s1)
            )
            db.commit()

    def test_individual_employee_attendance_doc(self):
        """Verify individual employee attendance document format, days 1 to 31, and contract 10h+2h OT."""
        with self.app.app_context():
            db = get_db()
            emp = db.execute(
                "SELECT e.*, s.name shift_name, c.name category_name, c.code category_code, c.work_hours category_work_hours "
                "FROM employees e LEFT JOIN shifts s ON s.id=e.shift_id "
                "LEFT JOIN employee_categories c ON c.id=e.category_id WHERE e.emp_code='CON001'"
            ).fetchone()

            buf = build_employee_attendance_doc(emp, date(2026, 10, 1), date(2026, 10, 31), db)
            self.assertIsNotNone(buf)

            # Read workbook back
            wb = openpyxl.load_workbook(buf)
            ws = wb.active
            self.assertEqual(ws.title, "CON001 Attendance")

            # Check Employee ID in document
            self.assertEqual(ws["B4"].value, "CON001")
            self.assertEqual(ws["E4"].value, "Contract Worker Ram")

            # Check Day 1 row (Row 9)
            self.assertEqual(ws["A9"].value, "Day 1")
            self.assertEqual(ws["D9"].value, "OT")
            self.assertEqual(ws["G9"].value, 10.0)  # Regular hours: 10 hrs
            self.assertEqual(ws["H9"].value, 2.0)   # Overtime: 2 hrs
            self.assertEqual(ws["I9"].value, 12.0)  # Total hours: 12 hrs
            self.assertIn("10 hrs present + 2.0 hrs OT present", ws["J9"].value)

            # Check Day 31 row (Row 39)
            self.assertEqual(ws["A39"].value, "Day 31")
            self.assertIn(ws["D39"].value, ("-", "A"))
            self.assertIn(ws["J39"].value, ("Not Checked In", "Absent (No punch recorded)"))

            # Check Total Days Come in summary (30 days)
            summary_text = [ws.cell(row=r, column=7).value for r in range(40, 50) if ws.cell(row=r, column=7).value]
            self.assertIn("30 days", summary_text)

    def test_hr_export_employee_endpoint(self):
        """Verify HR endpoint downloads document named {emp_code}_attendance_{month}.xlsx."""
        with self.app.app_context():
            db = get_db()
            admin = db.execute("SELECT * FROM employees WHERE role='hr'").fetchone()
            con_emp = db.execute("SELECT * FROM employees WHERE emp_code='CON001'").fetchone()

        with self.client.session_transaction() as sess:
            sess["uid"] = admin["id"]

        res = self.client.get(f"/hr/reports/employee/{con_emp['id']}/export?month=2026-10")
        self.assertEqual(res.status_code, 200)
        self.assertIn("attachment;", res.headers.get("Content-Disposition", ""))
        self.assertIn("CON001_attendance_2026_10.xlsx", res.headers.get("Content-Disposition", ""))

    def test_employee_self_export_endpoint(self):
        """Verify employee can download their own attendance document in app."""
        with self.app.app_context():
            db = get_db()
            con_emp = db.execute("SELECT * FROM employees WHERE emp_code='CON001'").fetchone()

        with self.client.session_transaction() as sess:
            sess["uid"] = con_emp["id"]

        res = self.client.get("/history/export?month=2026-10")
        self.assertEqual(res.status_code, 200)
        self.assertIn("attachment;", res.headers.get("Content-Disposition", ""))
        self.assertIn("CON001_attendance_2026_10.xlsx", res.headers.get("Content-Disposition", ""))

    def test_category_filtered_master_export(self):
        """Verify HR master export with category filter downloads {category_code}_attendance_{month}.xlsx."""
        with self.app.app_context():
            db = get_db()
            admin = db.execute("SELECT * FROM employees WHERE role='hr'").fetchone()
            cat = db.execute("SELECT * FROM employee_categories WHERE code='CONTRACT'").fetchone()

        with self.client.session_transaction() as sess:
            sess["uid"] = admin["id"]

        res = self.client.get(f"/hr/reports/export?month=2026-10&category={cat['id']}")
        self.assertEqual(res.status_code, 200)
        self.assertIn("CONTRACT_attendance_2026_10.xlsx", res.headers.get("Content-Disposition", ""))

        wb = openpyxl.load_workbook(io.BytesIO(res.data))
        sheet_names = wb.sheetnames
        self.assertIn("Summary", sheet_names)
        self.assertIn("Day-wise (Day 1 to 31)", sheet_names)
        self.assertIn("Contract Attendance & OT", sheet_names)


if __name__ == "__main__":
    unittest.main()
